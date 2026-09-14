#!/usr/bin/env python
# make_docx.py - собирает OTCHET.md в файл Word (OTCHET.docx) с вставленными картинками.
# Поддерживается тот набор разметки, который реально используется в отчёте:
# заголовки # ## ###, абзацы, **жирный**, `моноширинный`, таблицы, картинки,
# цитаты (>), маркированные (*) и нумерованные (1.) списки, разделители ---.
import io, os, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

H = os.path.dirname(os.path.abspath(__file__))
MD = os.path.join(H, "OTCHET.md")
OUT = os.path.join(H, "OTCHET.docx")
IMG_W = Cm(17.0)          # ширина картинки при полях 2 см

ACCENT = RGBColor(0x1F, 0x4E, 0x79)

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.0)
    s.top_margin = s.bottom_margin = Cm(2.0)
st = doc.styles["Normal"]
st.font.name = "Calibri"; st.font.size = Pt(10.5)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.12

for lvl, size in ((1, 18), (2, 14), (3, 12)):
    h = doc.styles["Heading %d" % lvl]
    h.font.name = "Calibri"; h.font.size = Pt(size)
    h.font.color.rgb = ACCENT; h.font.bold = True


def add_runs(par, text):
    """Разложить строку на обычный / **жирный** / `код` и добавить в абзац."""
    for piece in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            r = par.add_run(piece[2:-2]); r.bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            r = par.add_run(piece[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
        else:
            par.add_run(piece.replace("\\|", "|"))


def split_row(line):
    cells = line.strip().strip("|").split("|")
    return [c.strip() for c in cells]


def add_table(rows):
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j in range(ncol):
            txt = row[j] if j < len(row) else ""
            p = cells[j].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, txt)
            for r in p.runs:
                r.font.size = Pt(9)
                if i == 0:
                    r.bold = True
    doc.add_paragraph()


lines = io.open(MD, encoding="utf-8").read().replace("\r\n", "\n").split("\n")
i, buf, tbl = 0, [], []


def flush_par():
    global buf
    if buf:
        p = doc.add_paragraph()
        add_runs(p, " ".join(buf))
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        buf = []


def flush_tbl():
    global tbl
    if tbl:
        add_table(tbl)
        tbl = []


while i < len(lines):
    ln = lines[i].rstrip()
    s = ln.strip()

    m = re.match(r'^(#{1,3})\s+(.*)$', s)
    if m:
        flush_par(); flush_tbl()
        doc.add_heading(m.group(2), level=len(m.group(1)))
        i += 1; continue

    m = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)$', s)
    if m:
        flush_par(); flush_tbl()
        path = os.path.join(H, m.group(2).replace("/", os.sep))
        if os.path.exists(path):
            doc.add_picture(path, width=IMG_W)
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = doc.add_paragraph()
            r = cap.add_run(m.group(1)); r.italic = True; r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        i += 1; continue

    if s.startswith("|") and s.endswith("|"):
        flush_par()
        if re.match(r'^\|[\s:\-|]+\|$', s):     # строка-разделитель шапки
            i += 1; continue
        tbl.append(split_row(s))
        i += 1; continue
    flush_tbl()

    if s.startswith("---") and set(s) <= set("-"):
        flush_par()
        p = doc.add_paragraph(); p.add_run("─" * 60).font.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        i += 1; continue

    if s.startswith(">"):
        flush_par()
        quote = []
        while i < len(lines) and lines[i].strip().startswith(">"):
            quote.append(lines[i].strip().lstrip(">").strip()); i += 1
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(8)
        add_runs(p, " ".join(quote))
        for r in p.runs:
            r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
        continue

    m = re.match(r'^[*-]\s+(.*)$', s)
    if m:
        flush_par()
        item = [m.group(1)]; i += 1
        while i < len(lines) and lines[i].startswith("  ") and lines[i].strip() \
                and not re.match(r'^\s*[*-]\s', lines[i]) and not re.match(r'^\s*\d+\.\s', lines[i]):
            item.append(lines[i].strip()); i += 1
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, " ".join(item))
        continue

    m = re.match(r'^(\d+)\.\s+(.*)$', s)
    if m:
        flush_par()
        item = [m.group(2)]; i += 1
        while i < len(lines) and lines[i].startswith("   ") and lines[i].strip():
            item.append(lines[i].strip()); i += 1
        p = doc.add_paragraph(style="List Number")
        add_runs(p, " ".join(item))
        continue

    if not s:
        flush_par(); i += 1; continue

    buf.append(s); i += 1

flush_par(); flush_tbl()
doc.save(OUT)
print("->", OUT, "%.2f МБ" % (os.path.getsize(OUT) / 1048576.))
